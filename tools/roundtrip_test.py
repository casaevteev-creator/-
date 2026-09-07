# -*- coding: utf-8 -*-
"""Проверка нового формата ввода: канонические данные июня -> единый файл ->
обратное чтение -> расчёт. Итоги должны совпасть с уже сверённым отчётом за июнь."""
import datetime
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import input_xlsx  # noqa: E402
import model  # noqa: E402
from model import GROUPS, GROUP_RU  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent


def fill(canon, path):
    year, month = canon["meta"]["year"], canon["meta"]["month"]
    input_xlsx.make_template(path, year, month)
    wb = openpyxl.load_workbook(path)

    # --- 1 Оплаты: восстанавливаем из дневных данных, масштабируя карты к итогам групп
    ws = wb["1 Оплаты"]
    daily, bg = canon["revenue"]["daily"], canon["revenue"]["by_group"]
    scale = {"cash": {}, "card": {}}
    for way in ("cash", "card"):
        for g in GROUPS:
            s = sum(d[way].get(g, 0) for d in daily)
            scale[way][g] = (bg[g][way] / s) if s else 1.0
    r, pid = 3, 0
    for d in daily:
        date = datetime.date.fromisoformat(d["date"])
        for g in GROUPS:
            for way, amount in ((input_xlsx.PAY_CASH, d["cash"].get(g, 0) * scale["cash"][g]),
                                (input_xlsx.PAY_CARD, d["card"].get(g, 0) * scale["card"][g])):
                if not amount:
                    continue
                pid += 1
                ws.cell(row=r, column=1, value=f"P{pid:05d}")
                ws.cell(row=r, column=2, value=date)
                ws.cell(row=r, column=3, value="старая")
                ws.cell(row=r, column=6, value=GROUP_RU[g])
                ws.cell(row=r, column=8, value=round(amount, 2))
                ws.cell(row=r, column=9, value=way)
                r += 1
    last = datetime.date(year, month, daily[-1]["day"])
    for g in GROUPS:
        for way, amount, refund in ((input_xlsx.PAY_ACCT, bg[g]["bank_ind"], None),
                                    (input_xlsx.PAY_ACCT, bg[g]["refund"], "да")):
            if not amount:
                continue
            pid += 1
            ws.cell(row=r, column=1, value=f"P{pid:05d}")
            ws.cell(row=r, column=2, value=last)
            ws.cell(row=r, column=3, value="старая")
            ws.cell(row=r, column=6, value=GROUP_RU[g])
            ws.cell(row=r, column=8, value=round(amount, 2))
            ws.cell(row=r, column=9, value=way)
            ws.cell(row=r, column=10, value=refund)
            r += 1

    # --- 2 Счёт 51
    ws = wb["2 Счёт 51"]
    ws.cell(row=3, column=2, value=canon["cashflow"]["opening"])
    ws.cell(row=4, column=2, value=canon["cashflow"]["closing"])
    r = 5
    for f in canon["cashflow"]["flows"]:
        ws.cell(row=r, column=1, value=f["label"])
        ws.cell(row=r, column=2, value=f.get("debit"))
        ws.cell(row=r, column=3, value=f.get("credit"))
        ws.cell(row=r, column=4, value=f.get("note"))
        r += 1

    # --- 3 Счёт 60 (агрегаты по статьям — как их отдаст 1С с проставленным блоком)
    ws = wb["3 Счёт 60"]
    r = 3
    for block, items in (("Медицинские + ИП", canon["expenses"]["medical"]),
                         ("Управленческие", canon["expenses"]["management"])):
        for it in items:
            ws.cell(row=r, column=1, value=it["name"])
            ws.cell(row=r, column=2, value=it["amount"])
            ws.cell(row=r, column=4, value=block)
            ws.cell(row=r, column=5, value=it["name"])
            r += 1
    halva = canon["costs"].get("halva") or {}
    hv = (halva.get("mal") or 0) + (halva.get("pog") or 0)
    if hv:
        ws.cell(row=r, column=1, value="Халва (рассрочка)")
        ws.cell(row=r, column=2, value=hv)
        ws.cell(row=r, column=4, value="Управленческие")
        ws.cell(row=r, column=5, value="Халва (рассрочка)")

    # --- 4 Методика расходов
    ws = wb["4 Методика расходов"]
    keys = ["payroll_taxes", "salary", "cash_expenses", "bank_common",
            "bank_acquiring", "alimony", "adjust", "suppliers_bank"]
    for i, key in enumerate(keys, start=3):
        c = canon["costs"].get(key) or {}
        ws.cell(row=i, column=2, value=c.get("total"))
        ws.cell(row=i, column=3, value=c.get("mal"))
        ws.cell(row=i, column=4, value=c.get("pog"))

    # --- 5 Ручные данные
    ws = wb["5 Ручные данные"]
    ref, man = canon["founders_reference"], canon["manual"]
    vals = {
        "Остаток с прошлого месяца · Маланичев": ref["prev_balance"]["mal"],
        "Остаток с прошлого месяца · Погосян": ref["prev_balance"]["pog"],
        "Резерв на аренду · всего": man.get("reserve_total"),
        "Резерв на аренду · удержано в этом месяце · Маланичев": man["reserve_month"]["mal"],
        "Резерв на аренду · удержано в этом месяце · Погосян": man["reserve_month"]["pog"],
        "Долг по кредитной линии": man.get("credit_debt"),
        "Остаток наличных в кассе на конец месяца": man.get("cash_on_hand"),
        "Проценты по депозитам за месяц": man.get("deposit_interest"),
    }
    for row in range(3, ws.max_row + 1):
        k = ws.cell(row=row, column=1).value
        if k in vals:
            ws.cell(row=row, column=2, value=vals[k])

    # --- 6 История выручки
    ws = wb["6 История выручки"]
    for row in range(3, ws.max_row + 1):
        yy = ws.cell(row=row, column=1).value
        mm = ws.cell(row=row, column=2).value
        if yy is None or mm is None:
            continue
        mnum = model.MONTHS_RU.index(str(mm)) + 1
        v = canon["history"].get(str(int(yy)), {}).get(str(mnum))
        if v:
            ws.cell(row=row, column=3, value=v)
    wb.save(path)
    return path


def main():
    canon = json.loads((ROOT / "data" / "canonical" / "2026-06.json").read_text(encoding="utf-8"))
    ref = model.compute(canon)
    hv = ((canon["costs"].get("halva") or {}).get("mal") or 0) + \
         ((canon["costs"].get("halva") or {}).get("pog") or 0)
    tmp = ROOT / "out" / "roundtrip-Июнь-2026.xlsx"
    fill(canon, tmp)
    back = input_xlsx.read(tmp, 2026, 6)
    got = model.compute(back)

    pairs = [
        ("выручка", ref["revenue"]["total"], got["revenue"]["total"]),
        ("расходы сч.60 (+Халва)", ref["expenses"]["acc60_total"] + hv, got["expenses"]["acc60_total"]),
        ("доход учредителей", ref["founders_total_income"], got["founders_total_income"]),
        ("доход Маланичева", ref["founders"]["mal"]["income"], got["founders"]["mal"]["income"]),
        ("доход Погосяна", ref["founders"]["pog"]["income"], got["founders"]["pog"]["income"]),
        ("к выплате Маланичеву", ref["founders"]["mal"]["payout"], got["founders"]["mal"]["payout"]),
        ("к выплате Погосяну", ref["founders"]["pog"]["payout"], got["founders"]["pog"]["payout"]),
        ("итог водопада", ref["waterfall_result"], got["waterfall_result"]),
        ("остаток на счёте", ref["cashflow"]["closing"], got["cashflow"]["closing"]),
        ("рост с начала года, %", ref["trend"]["ytd_growth"], got["trend"]["ytd_growth"]),
    ]
    ok = True
    for name, a, b in pairs:
        d = abs((a or 0) - (b or 0))
        good = d <= 1.0
        ok &= good
        print(f"{'OK ' if good else 'ХХ '} {name:26s} эталон {a:>18,.2f}   новый формат {b:>18,.2f}   Δ {d:,.2f}")
    print("\nИТОГ:", "новый формат воспроизводит отчёт" if ok else "ЕСТЬ РАСХОЖДЕНИЯ")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
