import openpyxl, json
F="/root/.claude/uploads/ac5b8ad0-f2cc-56ed-bd75-26d657f80f0f/55319b99-________________2026.xlsx"
ws=openpyxl.load_workbook(F,data_only=True)["Контрагенты"]
book={}; arts={}; vend=[]
for r in range(6,125):
    n=ws.cell(row=r,column=2).value; p=ws.cell(row=r,column=3).value
    a=ws.cell(row=r,column=4).value; b=ws.cell(row=r,column=5).value
    if not n or not a: continue
    n=" ".join(str(n).split()); a=" ".join(str(a).split()); b=" ".join(str(b).split()) if b else ""
    book[n]=[a,b]; arts.setdefault(a,b)
    vend.append([n, round(float(p),2)])
vend.append(["КВАЛИТЕТ ООО", 11000000.0])
book["КВАЛИТЕТ ООО"]=["Ремонт Самокатная 1 стр.12","управленческие"]
arts.setdefault("Ремонт Самокатная 1 стр.12","управленческие")
BOOK={"vendors":book,"articles":arts}
EX={"vendors":vend,
 "groups":{"mal":3840000.0,"pog":11985000.0,"other":78426470.30,"cosm":2788990.0,"misc":298270.0},
 "params":{"tax":6135792.28,"sal":12490831.85,"bank":450077.22,"ali":38726.66,"int":204388.93,
   "corr":220355.37,"mAn":30620.0,"mMat":140241.64,"mAss":64884.23,"mAcq":15981.06,
   "pAn":12248.0,"pMat":1479282.06,"pAss":226384.0,"pAcq":49901.51,"dep":396337.24,
   "prevM":30658896.02,"prevP":45275972.21,"divM":5882353.0,"divP":5882353.0,
   "revBank":97338730.30,"refund":465000.0}}
html="\n".join(open('app/part%d.html'%i,encoding='utf-8').read() for i in (1,2,6,3,4,5))
html=html.replace("__BOOK__", json.dumps(BOOK,ensure_ascii=False,separators=(',',':')))
html=html.replace("__EXAMPLE__", json.dumps(EX,ensure_ascii=False,separators=(',',':')))
open('/home/user/-/out/upravlenka.html','w',encoding='utf-8').write(html)
print("собрано, символов:",len(html))
