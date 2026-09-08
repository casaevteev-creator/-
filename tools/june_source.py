# -*- coding: utf-8 -*-
"""Адаптер: три исходных файла бухгалтерии (июньский формат) -> канонический JSON.

Нужен для того, чтобы:
  1) доказать, что расчётная модель воспроизводит уже сверённый отчёт за июнь до рубля;
  2) иметь эталон, с которым сверяется новый (оптимизированный) формат выгрузок.
"""
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl
import xlrd

GROUPS = ["mal", "pog", "other", "misc", "cosm"]
GROUP_RU = {
    "mal": "Маланичев",
    "pog": "Погосян",
    "other": "Другие хирурги",
    "misc": "Прочие доходы",
    "cosm": "Косметология",
}
HEADER_TO_GROUP = {
    "маланичев": "mal",
    "погосян": "pog",
    "другие хирурги": "other",
    "др.хирурги": "other",
    "пр.доходы": "misc",
    "прочие доходы": "misc",
    "прочее": "misc",
    "косметология": "cosm",
}
MONTHS_RU = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
             "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
MONTH_NUM = {m.lower(): i + 1 for i, m in enumerate(MONTHS_RU)}


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def txt(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


# --------------------------------------------------------------------------- ВЫРУЧКА
def _sheet_for(wb, year, month):
    want = (MONTHS_RU[month - 1].lower(), str(year))
    for name in wb.sheetnames:
        m = re.match(r"\s*([А-Яа-яЁё]+)\s*\(?\d*\)?\s*(\d{4})", name.strip())
        if m and m.group(1).lower() == want[0] and m.group(2) == want[1]:
            return wb[name]
    raise KeyError(f"нет листа за {MONTHS_RU[month-1]} {year}")


def _revenue_layout(ws):
    """Находит строку шапки и колонки групп (пара «Касса / Банк» на каждую группу)."""
    hrow = tcol = None
    for row in ws.iter_rows(min_row=1, max_row=15, max_col=32):
        for c in row:
            if isinstance(c.value, str) and "Всего за день" in c.value:
                hrow, tcol = c.row, c.column
        if hrow:
            break
    if not hrow:
        raise ValueError("не найдена шапка «Всего за день»")

    cols = {}
    for col in range(1, tcol):
        label = txt(ws.cell(row=hrow - 1, column=col).value).lower()
        g = HEADER_TO_GROUP.get(label)
        if g:
            cols[g] = {"cash": col, "card": col + 1}
    return {
        "hrow": hrow, "total": tcol, "kkt": tcol + 1, "diff": tcol + 2, "groups": cols,
    }


def _find_row(ws, needle, maxcol=6, start=1):
    needle = needle.lower()
    for row in range(start, ws.max_row + 1):
        for col in range(1, maxcol + 1):
            v = ws.cell(row=row, column=col).value
            if isinstance(v, str) and needle in v.lower():
                return row, col
    return None, None


def _month_revenue(ws):
    """Значение строки «Выручка за месяц» — кассовая база, как в графике динамики."""
    row, col = _find_row(ws, "выручка", maxcol=3)
    while row:
        below = txt(ws.cell(row=row + 1, column=col).value).lower()
        if below.startswith("за месяц"):
            for cc in range(col + 1, col + 5):
                v = num(ws.cell(row=row, column=cc).value)
                if v:
                    return v
        row, col = _find_row(ws, "выручка", maxcol=3, start=row + 1)
    return None


def parse_revenue(path, year, month):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = _sheet_for(wb, year, month)
    lay = _revenue_layout(ws)

    daily, seen = [], set()
    for row in range(lay["hrow"] + 1, ws.max_row + 1):
        d = None
        for col in range(1, 4):
            v = ws.cell(row=row, column=col).value
            if isinstance(v, datetime.datetime):
                d = v
        if not d or d.year != year or d.month != month or d.day in seen:
            continue
        seen.add(d.day)
        rec = {"day": d.day, "date": d.strftime("%Y-%m-%d"), "cash": {}, "card": {}}
        for g, cc in lay["groups"].items():
            rec["cash"][g] = num(ws.cell(row=row, column=cc["cash"]).value) or 0.0
            rec["card"][g] = num(ws.cell(row=row, column=cc["card"]).value) or 0.0
        rec["total"] = num(ws.cell(row=row, column=lay["total"]).value) or 0.0
        rec["kkt"] = num(ws.cell(row=row, column=lay["kkt"]).value)
        rec["kkt_diff"] = num(ws.cell(row=row, column=lay["diff"]).value)
        daily.append(rec)
    daily.sort(key=lambda r: r["day"])

    # эквайринг «в отчёт» — с поправкой на стыковые дни месяца
    card_report, row = {}, _find_row(ws, "сумма в отчет", maxcol=3)[0]
    if row:
        for g, cc in lay["groups"].items():
            for probe in (cc["cash"], cc["card"]):
                v = num(ws.cell(row=row, column=probe).value)
                if v:
                    card_report[g] = v
                    break

    # комиссия эквайринга по группам (правый справочный блок)
    fee = {}
    for row in range(1, ws.max_row + 1):
        for col in range(1, 24):
            label = txt(ws.cell(row=row, column=col).value).lower().rstrip(".")
            key = {"маланичев": "mal", "погосян": "pog", "др.хирурги": "other",
                   "пр.доходы": "misc", "косметология": "cosm"}.get(label)
            if key and key not in fee:
                v = num(ws.cell(row=row, column=col + 1).value)
                if v and v < 1_000_000:
                    fee[key] = v
    return {
        "daily": daily,
        "cash_by_group": {g: round(sum(r["cash"][g] for r in daily), 2) for g in lay["groups"]},
        "card_by_day_by_group": {g: round(sum(r["card"][g] for r in daily), 2) for g in lay["groups"]},
        "card_report_by_group": card_report,
        "month_cash_basis": _month_revenue(ws),
        "acquiring_fee": fee,
        "acquiring_fee_total": round(sum(fee.values()), 2) if fee else None,
    }


def parse_history(path, years=(2025, 2026)):
    wb = openpyxl.load_workbook(path, data_only=True)
    hist = {str(y): {} for y in years}
    for name in wb.sheetnames:
        m = re.match(r"\s*([А-Яа-яЁё]+)\s*\(?\d*\)?\s*(\d{4})", name.strip())
        if not m:
            continue
        mon, yr = MONTH_NUM.get(m.group(1).lower()), int(m.group(2))
        if not mon or yr not in years:
            continue
        v = _month_revenue(wb[name])
        if v:
            hist[str(yr)][str(mon)] = round(v, 2)
    return hist


# --------------------------------------------------------------------------- БАНК (сч. 51)
def parse_bank(path):
    ws = openpyxl.load_workbook(path, data_only=True)["TDSheet"]

    def cell(r, c):
        return num(ws.cell(row=r, column=c).value)

    def row_of(needle, maxcol=4, start=1):
        return _find_row(ws, needle, maxcol=maxcol, start=start)[0]

    # --- сводная выручка (блок «Выручка / Маланичев / Погосян / ...»)
    head = row_of("выручка", maxcol=3)
    while head and txt(ws.cell(row=head, column=4).value).lower() != "маланичев":
        head = row_of("выручка", maxcol=3, start=head + 1)
    gcol = {}
    for col in range(3, 12):
        label = txt(ws.cell(row=head, column=col).value).lower()
        if label == "выручка":
            gcol["total"] = col
        elif label in HEADER_TO_GROUP:
            gcol[HEADER_TO_GROUP[label]] = col

    def block_row(needle, start=head):
        return _find_row(ws, needle, maxcol=3, start=start)[0]

    rows = {
        "bank_ind": block_row("оплата на р/с"),
        "card": block_row("банк"),
        "cash": block_row("касса"),
        "refund": block_row("возврат пациенту"),
        "total": block_row("итого выручка"),
    }
    revenue = {}
    for key, r in rows.items():
        revenue[key] = {k: (cell(r, c) or 0.0) for k, c in gcol.items()}

    # --- расходы и методика
    def line(needle, start=rows["total"]):
        r = _find_row(ws, needle, maxcol=3, start=start)[0]
        if not r:
            return None
        return {"row": r, "total": cell(r, 3), "mal": cell(r, 4), "pog": cell(r, 5)}

    costs = {
        "payroll_taxes": line("налоги зп"),
        "writ_payments": line("выплата по исполнит"),
        "salary": line("зарплата по банку"),
        "cash_expenses": line("раходы по кассе"),
        "bank_common": line("услуги банка (общие)"),
        "bank_acquiring": line("услуги банка (эквайринг)"),
        "alimony": line("алименты"),
        "credit_interest": line("% по кредиту"),
        "adjust": line("доп.расходы"),
        "suppliers_bank": line("поставщики банк"),
        "utilities": line("аренда коммун"),
        "halva": line("халва"),
        "ip_surgeons": line("ип хирурги"),
        "total": line("итого расходы"),
    }
    r_u = costs["utilities"]["row"]
    costs["rent"] = {"row": r_u + 1, "total": cell(r_u + 1, 3), "mal": cell(r_u + 1, 4), "pog": cell(r_u + 1, 5)}

    r_sup = _find_row(ws, "в т.ч. :             операцион", maxcol=3, start=rows["total"])[0]
    if r_sup:
        costs["operational"] = {"row": r_sup, "total": cell(r_sup, 3), "mal": cell(r_sup, 4), "pog": cell(r_sup, 5)}
        costs["common"] = {"row": r_sup + 1, "total": cell(r_sup + 1, 3),
                           "mal": cell(r_sup + 1, 4), "pog": cell(r_sup + 1, 5)}
    r_sb = costs["suppliers_bank"]["row"]
    costs["suppliers"] = {"row": r_sb + 5, "total": cell(r_sb + 5, 3),
                          "mal": cell(r_sb + 5, 4), "pog": cell(r_sb + 5, 5)}

    deposit_share = line("% по депозитам за")
    dividends = line("выплачены дивиденды")
    rent_paid = line("оплата аренды")
    reserve_return = None
    if rent_paid:
        rr = rent_paid["row"] + 1
        reserve_return = {"row": rr, "total": cell(rr, 3), "mal": cell(rr, 4), "pog": cell(rr, 5)}
    income = line("итого за месяц")
    prev = line("остаток с предыд")
    payout = line("итого к выплате")
    reserve_total_row = _find_row(ws, "отложено на аренду", maxcol=3)[0]
    reserve_month = line("из них в")

    def scalar(needle):
        r = _find_row(ws, needle, maxcol=3)[0]
        return cell(r, 3) if r else None

    # --- ДДС по корсчетам
    flows = []
    for r in range(1, 70):
        label = txt(ws.cell(row=r, column=2).value)
        note = txt(ws.cell(row=r, column=5).value) or txt(ws.cell(row=r, column=6).value)
        d, k = cell(r, 3), cell(r, 4)
        if label and (d or k):
            flows.append({"row": r, "label": label, "debit": d, "credit": k, "note": note})

    opening = _find_row(ws, "начальное сальдо", maxcol=3)[0]
    closing = _find_row(ws, "конечное сальдо", maxcol=3)[0]
    turn = _find_row(ws, "оборот", maxcol=3)[0]

    return {
        "revenue": revenue,
        "costs": {k: v for k, v in costs.items() if v},
        "income": income,
        "prev_balance": prev,
        "payout": payout,
        "deposit_share": deposit_share,
        "dividends": dividends,
        "reserve_return": reserve_return,
        "reserve_total": cell(reserve_total_row, 3) if reserve_total_row else None,
        "reserve_month": reserve_month,
        "credit_debt": scalar("долг по кредитной линии"),
        "cash_on_hand": scalar("остаток в кассе"),
        "deposit_interest": scalar("% по депозитам") or scalar("91.01 % по депозитам"),
        "credit_interest": None,
        "flows": flows,
        "opening": cell(opening, 3) if opening else None,
        "closing": cell(closing, 3) if closing else None,
        "turnover_debit": cell(turn, 3) if turn else None,
        "turnover_credit": cell(turn, 4) if turn else None,
    }


# --------------------------------------------------------------------------- РАСХОДЫ (сч. 60)
def parse_expenses(path):
    sh = xlrd.open_workbook(str(path)).sheet_by_index(0)
    vendors, medical, management, totals = [], [], [], {}
    bucket = None
    for i in range(sh.nrows):
        row = sh.row_values(i)
        name, amount, note = txt(row[0]), num(row[1]) if len(row) > 1 else None, txt(row[2]) if len(row) > 2 else ""
        if name and amount is not None and not name.lower().startswith("итого"):
            vendors.append({"name": name, "amount": round(amount, 2), "note": note})
        if name.lower().startswith("итого") and amount is not None:
            totals["vendors"] = round(amount, 2)

        cat, cval = txt(row[4]) if len(row) > 4 else "", num(row[5]) if len(row) > 5 else None
        low = cat.lower()
        if "медицинские расходы" in low:
            bucket = medical
            continue
        if "управленческие расходы" in low:
            bucket = management
            continue
        if low.startswith("итого") and bucket is not None:
            totals["medical" if bucket is medical else "management"] = round(cval, 2) if cval else None
            bucket = None
            continue
        if bucket is not None and cat and cval is not None:
            bucket.append({"name": cat, "amount": round(cval, 2)})
        if not cat and cval is not None and "grand" not in totals and totals.get("management"):
            totals["grand"] = round(cval, 2)
    return {"vendors": vendors, "medical": medical, "management": management, "totals": totals}


# --------------------------------------------------------------------------- сборка
def build(year, month, rev_path, bank_path, exp_path):
    rev = parse_revenue(rev_path, year, month)
    hist = parse_history(rev_path)
    bank = parse_bank(bank_path)
    exp = parse_expenses(exp_path)

    def g(block, key):
        return {k: round(v, 2) for k, v in bank["revenue"][block].items() if k == key or key is None}

    canon = {
        "meta": {
            "company": "ООО «ОМЕГА»", "brand": "Форма", "year": year, "month": month,
            "month_ru": MONTHS_RU[month - 1], "period": f"{MONTHS_RU[month-1]} {year}",
            "source": "три файла бухгалтерии (июньский формат)",
        },
        "revenue": {
            "daily": rev["daily"],
            "by_group": {
                grp: {
                    "cash": round(bank["revenue"]["cash"].get(grp, 0.0), 2),
                    "card": round(bank["revenue"]["card"].get(grp, 0.0), 2),
                    "bank_ind": round(bank["revenue"]["bank_ind"].get(grp, 0.0), 2),
                    "refund": round(bank["revenue"]["refund"].get(grp, 0.0), 2),
                    "total": round(bank["revenue"]["total"].get(grp, 0.0), 2),
                } for grp in GROUPS
            },
            "totals": {
                "cash": round(bank["revenue"]["cash"]["total"], 2),
                "card": round(bank["revenue"]["card"]["total"], 2),
                "bank_ind": round(bank["revenue"]["bank_ind"]["total"], 2),
                "refund": round(bank["revenue"]["refund"]["total"], 2),
                "total": round(bank["revenue"]["total"]["total"], 2),
            },
            "month_cash_basis": rev["month_cash_basis"],
            "acquiring_fee": rev["acquiring_fee"],
            "acquiring_fee_total": rev["acquiring_fee_total"],
            "kkt_diff_total": round(sum(r["kkt_diff"] or 0 for r in rev["daily"]), 2),
        },
        "history": hist,
        "expenses": exp,
        "cashflow": {
            "opening": bank["opening"], "closing": bank["closing"],
            "turnover_debit": bank["turnover_debit"], "turnover_credit": bank["turnover_credit"],
            "flows": bank["flows"],
        },
        "costs": bank["costs"],
        "founders_reference": {
            "income": bank["income"], "prev_balance": bank["prev_balance"], "payout": bank["payout"],
        },
        "manual": {
            "reserve_total": bank["reserve_total"],
            "reserve_month": bank["reserve_month"],
            "credit_debt": bank["credit_debt"],
            "cash_on_hand": bank["cash_on_hand"],
            "deposit_interest": bank["deposit_interest"],
        },
    }
    return canon


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    src = root / "data" / "source"
    canon = build(
        2026, 6,
        src / "Отчет-по-ВЫРУЧКЕ-ОМЕГА-июнь-2026.xlsx",
        src / "ОМЕГА-БАНК-июнь-2026.xlsx",
        src / "Расходы-июнь-2026.xls",
    )
    out = root / "data" / "canonical" / "2026-06.json"
    out.write_text(json.dumps(canon, ensure_ascii=False, indent=1), encoding="utf-8")
    print("записано:", out)
    r = canon["revenue"]["totals"]
    print("выручка:", r["total"], "| касса:", r["cash"], "| эквайринг:", r["card"])
    print("расходы 60:", canon["expenses"]["totals"])
    print("доход:", canon["founders_reference"]["income"])
