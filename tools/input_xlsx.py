# -*- coding: utf-8 -*-
"""Единый файл ввода вместо трёх старых экселей.

make_template() — создаёт пустой файл с шестью листами и подсказками.
read()          — читает заполненный файл в канонический словарь,
                  включая дедупликацию платежей в день перехода на новую управленку.
"""
import calendar
import datetime
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from model import GROUP_RU, GROUPS, MONTHS_RU

GROUP_BY_RU = {v: k for k, v in GROUP_RU.items()}
GROUP_BY_RU.update({"Др.хирурги": "other", "Прочее": "misc"})
PAY_CASH, PAY_CARD, PAY_ACCT = "Наличные", "Карта", "На р/с"

SHEETS = {
    "1 Оплаты": ["ID платежа", "Дата", "Система", "Пациент", "Врач", "Группа",
                 "Направление", "Сумма", "Способ оплаты", "Возврат", "ККТ"],
    "2 Счёт 51": ["Статья", "Поступление", "Списание", "Комментарий"],
    "3 Счёт 60": ["Контрагент", "Сумма", "Назначение", "Блок", "Статья"],
    "4 Методика расходов": ["Статья", "Всего", "Маланичев", "Погосян", "Правило"],
    "5 Ручные данные": ["Параметр", "Значение", "Комментарий"],
    "6 История выручки": ["Год", "Месяц", "Выручка"],
}
HINTS = {
    "1 Оплаты": "Одна строка = один платёж из управленки. «Система» — старая/новая. "
                "«Возврат» — «да» для возврата (сумма положительная). Группа: "
                + " / ".join(GROUP_RU[g] for g in GROUPS),
    "2 Счёт 51": "Анализ счёта 51. Строки «Начальное сальдо» и «Конечное сальдо» обязательны.",
    "3 Счёт 60": "Обороты счёта 60. «Блок»: «Медицинские + ИП» или «Управленческие». "
                 "«Статья» — укрупнённая категория для отчёта.",
    "4 Методика расходов": "Расходы, которые делятся между учредителями: ФОТ, налоги, услуги банка, "
                           "эквайринг, алименты, поставщики. «Правило» — справочно (50/50 или персонально).",
    "5 Ручные данные": "Цифры, которых нет в выгрузках — их диктует бухгалтерия.",
    "6 История выручки": "Помесячная выручка за текущий и прошлый год — для графика динамики.",
}
MANUAL_KEYS = [
    ("Остаток с прошлого месяца · Маланичев", None, "к выплате на конец прошлого месяца"),
    ("Остаток с прошлого месяца · Погосян", None, ""),
    ("Резерв на аренду · всего", None, "сколько всего отложено на депозит"),
    ("Резерв на аренду · удержано в этом месяце · Маланичев", None, ""),
    ("Резерв на аренду · удержано в этом месяце · Погосян", None, ""),
    ("Долг по кредитной линии", None, "на дату отчёта"),
    ("Остаток наличных в кассе на конец месяца", None, ""),
    ("Проценты по депозитам за месяц", None, ""),
    ("Дата перехода на новую управленку", None, "если перехода не было — оставить пустым"),
    ("Приоритетная система в день перехода", "новая", "какая система считается верной при дубле"),
]
METHOD_LINES = [
    ("Налоги с ФОТ (68.90 + 69)", "50/50"),
    ("Зарплата по банку", "по принадлежности сотрудников"),
    ("Расходы по кассе", "50/50"),
    ("Услуги банка (общие)", "50/50"),
    ("Услуги банка (эквайринг)", "персонально по своей выручке"),
    ("Алименты (удержание из ЗП)", "50/50"),
    ("Дополнительные корректировки", "ручная корректировка"),
    ("Поставщики (счёт 60)", "½ общих + персональные операционные"),
]

H_FILL = PatternFill("solid", fgColor="1C1B18")
H_FONT = Font(color="FFFFFF", bold=True, size=10)


def make_template(path, year, month):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        ws["A1"] = HINTS[name]
        ws["A1"].font = Font(italic=True, size=10, color="807A70")
        ws["A1"].alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(cols), 4))
        ws.row_dimensions[1].height = 32
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(row=2, column=i, value=c)
            cell.fill, cell.font = H_FILL, H_FONT
            ws.column_dimensions[get_column_letter(i)].width = max(14, min(38, len(c) + 10))
        ws.freeze_panes = ws.cell(row=3, column=1)

    ws = wb["5 Ручные данные"]
    for i, (k, v, note) in enumerate(MANUAL_KEYS, start=3):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
        ws.cell(row=i, column=3, value=note)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["C"].width = 48

    ws = wb["4 Методика расходов"]
    for i, (k, rule) in enumerate(METHOD_LINES, start=3):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=5, value=rule)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["E"].width = 40

    ws = wb["6 История выручки"]
    r = 3
    for yy in (year - 1, year):
        for mm in range(1, 13):
            if yy == year and mm > month:
                break
            ws.cell(row=r, column=1, value=yy)
            ws.cell(row=r, column=2, value=MONTHS_RU[mm - 1])
            r += 1

    ws = wb["2 Счёт 51"]
    for i, name in enumerate(["Начальное сальдо", "Конечное сальдо"], start=3):
        ws.cell(row=i, column=1, value=name)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["D"].width = 52

    wb.save(path)
    return path


def _rows(ws, cols):
    idx = {}
    for i, c in enumerate(cols, start=1):
        idx[c] = i
    for r in range(3, ws.max_row + 1):
        vals = {c: ws.cell(row=r, column=i).value for c, i in idx.items()}
        if any(v not in (None, "") for v in vals.values()):
            yield r, vals


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


def read(path, year, month, company="ООО «ОМЕГА»", brand="Форма"):
    wb = openpyxl.load_workbook(path, data_only=True)

    # ------------------------------------------------------- ручные данные
    manual_raw = {}
    for _, row in _rows(wb["5 Ручные данные"], SHEETS["5 Ручные данные"]):
        manual_raw[str(row["Параметр"]).strip()] = row["Значение"]

    def mv(key):
        return _num(manual_raw.get(key))

    switch_date = _date(manual_raw.get("Дата перехода на новую управленку"))
    priority = str(manual_raw.get("Приоритетная система в день перехода") or "новая").strip().lower()

    # ------------------------------------------------------- оплаты + дедуп
    payments, dropped = [], []
    for r, row in _rows(wb["1 Оплаты"], SHEETS["1 Оплаты"]):
        d = _date(row["Дата"])
        amount = _num(row["Сумма"])
        if d is None or amount is None:
            continue
        payments.append({
            "row": r, "id": str(row["ID платежа"] or "").strip(), "date": d,
            "system": str(row["Система"] or "").strip().lower(),
            "patient": str(row["Пациент"] or "").strip(),
            "doctor": str(row["Врач"] or "").strip(),
            "group": GROUP_BY_RU.get(str(row["Группа"] or "").strip(), "misc"),
            "way": str(row["Способ оплаты"] or "").strip(),
            "amount": amount,
            "refund": str(row["Возврат"] or "").strip().lower() in ("да", "yes", "1", "true"),
        })

    kept = []
    if switch_date:
        seen_id, seen_key = {}, {}
        same_day = [p for p in payments if p["date"] == switch_date]
        others = [p for p in payments if p["date"] != switch_date]
        preferred = "нов" if priority.startswith("нов") else "стар"
        same_day.sort(key=lambda p: 0 if p["system"].startswith(preferred) else 1)
        for p in same_day:
            key_id = p["id"] or None
            key = (p["patient"], p["doctor"], round(p["amount"], 2), p["way"], p["refund"])
            dup = (key_id and key_id in seen_id) or (key in seen_key)
            if dup:
                dropped.append({**p, "reason": "дубль в день перехода: платёж есть в обеих системах"})
                continue
            if key_id:
                seen_id[key_id] = True
            seen_key[key] = True
            kept.append(p)
        kept += others
    else:
        kept = payments

    daily_map = defaultdict(lambda: {"cash": defaultdict(float), "card": defaultdict(float)})
    by_group = {g: {"cash": 0.0, "card": 0.0, "bank_ind": 0.0, "refund": 0.0, "total": 0.0} for g in GROUPS}
    for p in kept:
        g = p["group"]
        if p["refund"]:
            by_group[g]["refund"] += p["amount"]
            continue
        if p["way"].lower().startswith("нал"):
            by_group[g]["cash"] += p["amount"]
            daily_map[p["date"]]["cash"][g] += p["amount"]
        elif p["way"].lower().startswith("карт"):
            by_group[g]["card"] += p["amount"]
            daily_map[p["date"]]["card"][g] += p["amount"]
        else:
            by_group[g]["bank_ind"] += p["amount"]
    for g in GROUPS:
        b = by_group[g]
        b["total"] = round(b["cash"] + b["card"] + b["bank_ind"] - b["refund"], 2)
        for k in b:
            b[k] = round(b[k], 2)

    days = calendar.monthrange(year, month)[1]
    daily = []
    for d in range(1, days + 1):
        date = datetime.date(year, month, d)
        rec = daily_map.get(date, {"cash": {}, "card": {}})
        cash = {g: round(rec["cash"].get(g, 0.0), 2) for g in GROUPS}
        card = {g: round(rec["card"].get(g, 0.0), 2) for g in GROUPS}
        daily.append({"day": d, "date": date.isoformat(), "cash": cash, "card": card,
                      "total": round(sum(cash.values()) + sum(card.values()), 2),
                      "kkt": None, "kkt_diff": None})

    totals = {k: round(sum(by_group[g][k] for g in GROUPS), 2)
              for k in ("cash", "card", "bank_ind", "refund", "total")}

    # ------------------------------------------------------- счёт 51
    flows, opening, closing = [], None, None
    for _, row in _rows(wb["2 Счёт 51"], SHEETS["2 Счёт 51"]):
        label = str(row["Статья"] or "").strip()
        d, k = _num(row["Поступление"]), _num(row["Списание"])
        low = label.lower()
        if "начальное сальдо" in low:
            opening = d if d is not None else k
            continue
        if "конечное сальдо" in low:
            closing = d if d is not None else k
            continue
        if d or k:
            flows.append({"label": label, "debit": d, "credit": k,
                          "note": str(row["Комментарий"] or "")})

    # ------------------------------------------------------- счёт 60
    vendors, blocks = [], {"Медицинские + ИП": defaultdict(float), "Управленческие": defaultdict(float)}
    for _, row in _rows(wb["3 Счёт 60"], SHEETS["3 Счёт 60"]):
        amount = _num(row["Сумма"])
        if amount is None:
            continue
        vendors.append({"name": str(row["Контрагент"] or "").strip(), "amount": round(amount, 2),
                        "note": str(row["Назначение"] or "")})
        block = str(row["Блок"] or "").strip()
        block = "Медицинские + ИП" if block.lower().startswith("мед") else "Управленческие"
        item = str(row["Статья"] or row["Назначение"] or "Прочее").strip()
        blocks[block][item] += amount
    medical = [{"name": k, "amount": round(v, 2)} for k, v in
               sorted(blocks["Медицинские + ИП"].items(), key=lambda x: -x[1])]
    management = [{"name": k, "amount": round(v, 2)} for k, v in
                  sorted(blocks["Управленческие"].items(), key=lambda x: -x[1])]
    med_total = round(sum(i["amount"] for i in medical), 2)
    mgmt_total = round(sum(i["amount"] for i in management), 2)

    # ------------------------------------------------------- методика расходов
    key_by_label = {
        "налоги с фот": "payroll_taxes", "зарплата": "salary", "расходы по кассе": "cash_expenses",
        "услуги банка (общие)": "bank_common", "услуги банка (эквайринг)": "bank_acquiring",
        "алименты": "alimony", "дополнительные корректировки": "adjust",
        "поставщики": "suppliers_bank",
    }
    costs = {}
    for r, row in _rows(wb["4 Методика расходов"], SHEETS["4 Методика расходов"]):
        label = str(row["Статья"] or "").strip().lower()
        key = next((v for k, v in key_by_label.items() if label.startswith(k)), None)
        if not key:
            continue
        mal, pog = _num(row["Маланичев"]) or 0.0, _num(row["Погосян"]) or 0.0
        costs[key] = {"row": r, "total": _num(row["Всего"]) if _num(row["Всего"]) is not None else mal + pog,
                      "mal": mal, "pog": pog}
    costs["total"] = {"row": None,
                      "total": round(sum(v["total"] or 0 for v in costs.values()), 2),
                      "mal": round(sum(v["mal"] for v in costs.values()), 2),
                      "pog": round(sum(v["pog"] for v in costs.values()), 2)}
    ip_line = next((i for i in medical if "ип хирург" in i["name"].lower()), None)
    costs["ip_surgeons"] = {"row": None, "total": ip_line["amount"] if ip_line else 0.0,
                            "mal": (ip_line["amount"] / 2 if ip_line else 0.0),
                            "pog": (ip_line["amount"] / 2 if ip_line else 0.0)}

    # ------------------------------------------------------- история
    history = {str(year - 1): {}, str(year): {}}
    for _, row in _rows(wb["6 История выручки"], SHEETS["6 История выручки"]):
        yy, mm, v = row["Год"], row["Месяц"], _num(row["Выручка"])
        if not v or yy is None:
            continue
        mm_num = (MONTHS_RU.index(str(mm).strip().capitalize()) + 1
                  if str(mm).strip().capitalize() in MONTHS_RU else _num(mm))
        if mm_num and str(int(yy)) in history:
            history[str(int(yy))][str(int(mm_num))] = round(v, 2)

    return {
        "meta": {"company": company, "brand": brand, "year": year, "month": month,
                 "month_ru": MONTHS_RU[month - 1], "period": f"{MONTHS_RU[month-1]} {year}",
                 "source": "единый файл ввода"},
        "revenue": {
            "daily": daily, "by_group": by_group, "totals": totals,
            "month_cash_basis": totals["cash"] + totals["card"],
            "acquiring_fee": {}, "acquiring_fee_total": None, "kkt_diff_total": None,
        },
        "history": history,
        "expenses": {"vendors": vendors, "medical": medical, "management": management,
                     "totals": {"medical": med_total, "management": mgmt_total,
                                "grand": round(med_total + mgmt_total, 2)}},
        "cashflow": {"opening": opening, "closing": closing,
                     "turnover_debit": round(sum(f["debit"] or 0 for f in flows), 2),
                     "turnover_credit": round(sum(f["credit"] or 0 for f in flows), 2),
                     "flows": flows},
        "costs": costs,
        "founders_reference": {
            "prev_balance": {"row": None,
                             "mal": mv("Остаток с прошлого месяца · Маланичев"),
                             "pog": mv("Остаток с прошлого месяца · Погосян")},
        },
        "manual": {
            "reserve_total": mv("Резерв на аренду · всего"),
            "reserve_month": {"row": None,
                              "mal": mv("Резерв на аренду · удержано в этом месяце · Маланичев") or 0.0,
                              "pog": mv("Резерв на аренду · удержано в этом месяце · Погосян") or 0.0},
            "credit_debt": mv("Долг по кредитной линии"),
            "cash_on_hand": mv("Остаток наличных в кассе на конец месяца"),
            "deposit_interest": mv("Проценты по депозитам за месяц"),
        },
        "migration": {
            "switch_date": switch_date.isoformat() if switch_date else None,
            "priority": priority,
            "payments_total": len(payments),
            "payments_kept": len(kept),
            "dropped": [{**d, "date": d["date"].isoformat()} for d in dropped],
        },
    }
